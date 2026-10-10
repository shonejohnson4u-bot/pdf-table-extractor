"""Local browser app for extracting user-selected PDF table layouts.

Install:
    py -m pip install opencv-python PyMuPDF numpy pytesseract openpyxl
Run:
    py .\guided_pdf_table_extractor.py

Tesseract OCR must be installed separately. Set TESSERACT_CMD if it is not on PATH.
"""

import io
import json
import os
import re
import shutil
import tempfile
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

import cv2
import fitz
import numpy as np
import openpyxl
import pytesseract
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HOST = "127.0.0.1"
PORT = 8765
MAX_UPLOAD = 150 * 1024 * 1024
RENDER_DPI = 180
SESSIONS = {}

TESSERACT_CMD = os.environ.get("TESSERACT_CMD", "")
if not TESSERACT_CMD:
    default_tesseract = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    TESSERACT_CMD = str(default_tesseract) if default_tesseract.exists() else (shutil.which("tesseract") or "")
if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD

HTML = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Guided PDF Table Extractor</title>
<style>
body{font:15px system-ui,Segoe UI,sans-serif;margin:0;background:#f3f5f8;color:#1c2530}
header{background:#18334f;color:#fff;padding:16px 24px;font-size:21px;font-weight:650}
main{max-width:1250px;margin:20px auto;padding:0 16px}
section{background:#fff;border:1px solid #d8dee7;border-radius:8px;padding:16px;margin:14px 0}
h2{font-size:18px;margin:0 0 12px}.row{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin:8px 0}
button{background:#1769aa;color:#fff;border:0;border-radius:5px;padding:9px 13px;cursor:pointer}
button.secondary{background:#536579}button:disabled{opacity:.45}
input{padding:8px;border:1px solid #adb8c5;border-radius:4px}
input[type=text]{min-width:245px}input[type=number]{width:85px}
#preview{position:relative;display:inline-block;max-width:100%;line-height:0;cursor:crosshair}
#pageImage{display:block;max-width:100%;height:auto}
#overlay{position:absolute;inset:0;width:100%;height:100%;pointer-events:auto}
.muted{color:#647386;font-size:13px}#status{white-space:pre-wrap;margin-top:10px}
.error{color:#a33}
</style></head><body>
<header>Guided PDF Table Extractor</header><main>
<section><h2>1. Upload PDF</h2>
<div class="row"><input id="pdf" type="file" accept="application/pdf">
<button id="upload">Upload PDF</button></div><div id="status" class="muted">No PDF loaded.</div></section>
<section id="editor" hidden><h2>2. Mark a reference table layout</h2>
<p class="muted">Choose a parts-list page (not a drawing). Drag over the body rows, starting below the heading row. Add column headings in left-to-right order. Divider marks are optional; without them, columns are divided evenly. Assign only pages with this table layout. Pages without row borders will be skipped and listed in the workbook.</p>
<div class="row"><button class="secondary" id="prev">Previous</button>
<label>Page <input id="page" type="number" min="1" value="1"></label>
<button class="secondary" id="next">Next</button><span id="pageCount" class="muted"></span></div>
<div id="preview"><img id="pageImage" alt="PDF page"><canvas id="overlay"></canvas></div>
<div class="row"><button class="secondary" id="clear">Clear marks</button><span id="markState" class="muted">Drag over the table body.</span></div>
<div class="row"><label>Layout name <input id="layoutName" type="text" placeholder="Parts list A"></label>
<label>Column headings <input id="headings" type="text" placeholder="Item no, Qty, Designation"></label></div>
<div class="row"><label>Pages from <input id="start" type="number" min="1"></label>
<label>through <input id="end" type="number" min="1"></label>
<button id="save">Save layout and page range</button></div>
<ul id="layouts"></ul><div class="row"><button id="export">Extract and download Excel</button></div>
</section></main>
<script>
const $=id=>document.getElementById(id);
let token="",total=0,pw=0,ph=0,box=null,dividers=[],drag=null;
function status(s,error=false){$("status").textContent=s;$("status").className=error?"error":""}
function current(){return Math.max(1,Math.min(total,Number($("page").value)||1))}
function paint(){
 const img=$("pageImage"),c=$("overlay");if(!img.naturalWidth)return;
 const r=img.getBoundingClientRect(),w=Math.round(r.width*devicePixelRatio),h=Math.round(r.height*devicePixelRatio);
 if(c.width!==w)c.width=w;if(c.height!==h)c.height=h;
 c.style.width=r.width+"px";c.style.height=r.height+"px";
 const x=c.getContext("2d");x.setTransform(c.width/pw,0,0,c.height/ph,0,0);x.clearRect(0,0,pw,ph);
 if(box){x.fillStyle="rgba(0,163,90,.10)";x.fillRect(box.x,box.y,box.w,box.h);
  x.strokeStyle="#00a35a";x.lineWidth=Math.max(3,pw/500);x.strokeRect(box.x,box.y,box.w,box.h);
  x.strokeStyle="#e04725";dividers.forEach(d=>{x.beginPath();x.moveTo(d,box.y);x.lineTo(d,box.y+box.h);x.stroke()})}
 if(drag){x.strokeStyle="#1769aa";x.lineWidth=Math.max(3,pw/500);x.strokeRect(drag.x,drag.y,drag.w,drag.h)}
}
function pagePoint(e){const r=$("pageImage").getBoundingClientRect();return{
 x:Math.max(0,Math.min(pw,(e.clientX-r.left)*pw/r.width)),
 y:Math.max(0,Math.min(ph,(e.clientY-r.top)*ph/r.height))}}
function loadPage(n){if(!token)return;n=Math.max(1,Math.min(total,n));$("page").value=n;
 const im=$("pageImage");im.onload=()=>{pw=im.naturalWidth;ph=im.naturalHeight;paint()};
 im.src=`/page?token=${encodeURIComponent(token)}&n=${n}&v=${Date.now()}`}
const canvas=$("overlay");
canvas.addEventListener("pointerdown",e=>{
 if(!token)return;const p=pagePoint(e);
 if(box&&p.x>box.x&&p.x<box.x+box.w&&p.y>=box.y&&p.y<=box.y+box.h){
  const d=dividers.find(v=>Math.abs(v-p.x)<pw*.012);
  if(d!==undefined){dividers=dividers.filter(v=>v!==d);$("markState").textContent="Divider removed.";paint();return}
  dividers.push(p.x);dividers.sort((a,b)=>a-b);$("markState").textContent=`${dividers.length} optional divider(s); click red divider to remove.`;paint();return}
 drag={x:p.x,y:p.y,w:0,h:0};canvas.setPointerCapture(e.pointerId);paint()});
canvas.addEventListener("pointermove",e=>{if(drag){const p=pagePoint(e);drag.w=p.x-drag.x;drag.h=p.y-drag.y;paint()}});
canvas.addEventListener("pointerup",()=>{if(!drag)return;let b={...drag};drag=null;
 if(b.w<0){b.x+=b.w;b.w=-b.w}if(b.h<0){b.y+=b.h;b.h=-b.h}
 if(b.w>10&&b.h>10){box=b;dividers=[];$("markState").textContent="Table body marked; divider marks are optional."}paint()});
$("upload").onclick=async()=>{const f=$("pdf").files[0];if(!f){status("Choose a PDF first.",true);return}
 if(f.size>150*1024*1024){status("PDF exceeds the 150 MB upload limit.",true);return}
 const fd=new FormData();fd.append("pdf",f);status("Uploading PDF…");
 try{const r=await fetch("/upload",{method:"POST",body:fd}),d=await r.json();if(!r.ok)throw Error(d.error);
 token=d.token;total=d.pages;$("pageCount").textContent=`of ${total}`;
 $("start").max=total;$("end").max=total;$("start").value=1;$("end").value=total;
 $("editor").hidden=false;box=null;dividers=[];status(`Loaded ${d.name} (${total} pages). Only assigned pages will be processed.`);loadPage(1)
 }catch(e){status(e.message,true)}};
$("prev").onclick=()=>loadPage(current()-1);$("next").onclick=()=>loadPage(current()+1);
$("page").onchange=()=>loadPage(current());window.addEventListener("resize",paint);
$("clear").onclick=()=>{box=null;dividers=[];paint();$("markState").textContent="Drag over the table body."};
function esc(s){return s.replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]))}
$("save").onclick=async()=>{if(!box){status("Mark the table body first.",true);return}
 const headings=$("headings").value.split(",").map(x=>x.trim()).filter(Boolean);
 const name=$("layoutName").value.trim(),start=Number($("start").value),end=Number($("end").value);
 if(!name||!headings.length){status("Enter a layout name and at least one heading. Divider marks are optional.",true);return}
 if(dividers.length>headings.length-1){status(`At most ${headings.length-1} divider(s) can be marked for ${headings.length} headings.`,true);return}
 if(!start||!end||start>end||start<1||end>total){status("Enter a valid page range.",true);return}
 const ds=[...dividers].sort((a,b)=>a-b);
 const item={name,headings,start,end,rect:[box.x/pw,box.y/ph,box.w/pw,box.h/ph],dividers:ds.map(x=>(x-box.x)/box.w)};
 try{const r=await fetch("/layout?token="+encodeURIComponent(token),{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(item)});
 const d=await r.json();if(!r.ok)throw Error(d.error);
 $("layouts").innerHTML=d.layouts.map(t=>`<li>${esc(t.name)} — pages ${t.start}–${t.end}</li>`).join("");
 status("Layout saved. Add other layouts for different formats.")
 }catch(e){status(e.message,true)}};
$("export").onclick=async()=>{if(!token){status("Upload a PDF first.",true);return}
 $("export").disabled=true;status("Extracting assigned pages; pages without table borders will be skipped…");
 try{const r=await fetch("/export?token="+encodeURIComponent(token),{method:"POST"});
 if(!r.ok){const d=await r.json();throw Error(d.error)}
 const blob=await r.blob(),url=URL.createObjectURL(blob),a=document.createElement("a");
 a.href=url;a.download=r.headers.get("X-Filename")||"parts.xlsx";a.click();URL.revokeObjectURL(url);
 status("Excel downloaded. Any pages without row borders are listed on the Skipped Pages sheet.")
 }catch(e){status(e.message,true)}finally{$("export").disabled=false}};
</script></body></html>"""


class PageSkipped(Exception):
    """Raised when a page has no usable row-border evidence for the assigned layout."""


def json_reply(handler, payload, status=200):
    body = json.dumps(payload).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def render(session, page_number, dpi=RENDER_DPI):
    with fitz.open(session["path"]) as doc:
        if not 1 <= page_number <= doc.page_count:
            raise ValueError("Page number is outside the PDF.")
        pix = doc[page_number - 1].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY, alpha=False)
        return pix.samples, pix.width, pix.height


def horizontal_rules(gray):
    """Detect long horizontal borders inside the manually selected table body."""
    height, width = gray.shape
    ink = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (max(20, int(width * 0.45)), 1))
    mask = cv2.morphologyEx(ink, cv2.MORPH_OPEN, kernel)
    positions = np.flatnonzero(np.count_nonzero(mask, axis=1) >= width * 0.45)
    if not len(positions):
        return []
    gap = max(3, int(height * 0.004))
    groups = np.split(positions, np.where(np.diff(positions) > gap)[0] + 1)
    return [int(round(float(np.mean(g)))) for g in groups if len(g)]


def infer_edges(width, divider_fractions, column_count):
    """Use optional marks as anchors and interpolate any unmarked boundaries."""
    anchors = [(0.0, 0), (1.0, column_count)]
    for fraction in sorted(divider_fractions):
        slot = max(1, min(column_count - 1, int(round(fraction * column_count))))
        anchors.append((fraction, slot))
    anchors.sort()
    fixed = []
    for anchor in anchors:
        if fixed and (anchor[0] <= fixed[-1][0] or anchor[1] <= fixed[-1][1]):
            raise ValueError("Divider marks do not form left-to-right column boundaries.")
        fixed.append(anchor)

    boundaries = [0] * (column_count + 1)
    boundaries[-1] = width
    for fraction, slot in fixed[1:-1]:
        boundaries[slot] = int(round(fraction * width))
    for (x0, c0), (x1, c1) in zip(fixed, fixed[1:]):
        for col in range(c0 + 1, c1):
            portion = (col - c0) / (c1 - c0)
            boundaries[col] = int(round((x0 + portion * (x1 - x0)) * width))
    if any(right <= left for left, right in zip(boundaries, boundaries[1:])):
        raise ValueError("Inferred columns are too narrow; check headings or divider marks.")
    return boundaries


def extract_page(session, page_number, layout):
    samples, width, height = render(session, page_number)
    image = np.frombuffer(samples, dtype=np.uint8).reshape(height, width).copy()
    rx, ry, rw, rh = layout["rect"]
    x0, y0 = round(rx * width), round(ry * height)
    x1, y1 = round((rx + rw) * width), round((ry + rh) * height)
    x0, x1, y0, y1 = max(0, x0), min(width, x1), max(0, y0), min(height, y1)
    if x1 - x0 < 20 or y1 - y0 < 20:
        raise PageSkipped("Selected table area is too small.")
    rules = horizontal_rules(image[y0:y1, x0:x1])
    if len(rules) < 2:
        raise PageSkipped("Could not find two horizontal row borders in the marked table area.")
    offsets = infer_edges(x1 - x0, layout["dividers"], len(layout["headings"]))
    edges = [x0 + offset for offset in offsets]
    rows = []
    for top, bottom in zip(rules, rules[1:]):
        if bottom - top < max(5, int(height * 0.003)):
            continue
        row = []
        for left, right in zip(edges, edges[1:]):
            px, py = max(2, int(width * 0.002)), max(2, int(height * 0.0015))
            cell = image[max(y0, y0 + top + py):min(y1, y0 + bottom - py),
                         min(x1, left + px):max(x0, right - px)]
            if cell.size == 0 or np.count_nonzero(cell < 170) < 5:
                row.append(("", 100))
                continue
            enlarged = cv2.resize(cell, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
            data = pytesseract.image_to_data(
                enlarged, config="--psm 6", lang="eng",
                output_type=pytesseract.Output.DICT
            )
            words = [(t.strip(), float(c)) for t, c in zip(data["text"], data["conf"])
                     if t.strip() and float(c) >= 0]
            if words:
                text = re.sub(r"\s+", " ", " ".join(t for t, _ in words))
                confidence = sum(c for _, c in words) / len(words)
                row.append((text, confidence))
            else:
                row.append(("?", 0))
        if any(text for text, _ in row):
            rows.append(row)
    return rows


def create_workbook(session):
    layouts = session["layouts"]
    if not layouts:
        raise ValueError("Save at least one layout and page range before exporting.")
    owners = {}
    for layout in layouts:
        for page in range(layout["start"], layout["end"] + 1):
            if page in owners:
                raise ValueError(f"Page {page} is assigned to multiple layouts.")
            owners[page] = layout["name"]

    book = openpyxl.Workbook()
    book.remove(book.active)
    skipped = []
    used = set()
    for layout in layouts:
        base = re.sub(r"[\[\]:*?/\\]", "-", layout["name"]).strip()[:31] or "Parts"
        name, index = base, 2
        while name.casefold() in used:
            suffix = f" ({index})"
            name, index = base[:31 - len(suffix)] + suffix, index + 1
        used.add(name.casefold())
        sheet = book.create_sheet(name)
        ncols = len(layout["headings"])
        if ncols > 1:
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        sheet.cell(1, 1, layout["name"]).font = Font(bold=True, size=14)
        sheet.cell(1, 1).alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 24
        for col, heading in enumerate(layout["headings"], 1):
            cell = sheet.cell(3, col, heading)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")

        for page in range(layout["start"], layout["end"] + 1):
            try:
                rows = extract_page(session, page, layout)
            except PageSkipped as exc:
                skipped.append((page, layout["name"], str(exc)))
                continue
            for row in rows:
                values = (row + [("", 100)] * ncols)[:ncols]
                sheet.append([text for text, _ in values])
                for col, (text, confidence) in enumerate(values, 1):
                    if text and confidence < 70:
                        sheet.cell(sheet.max_row, col).fill = PatternFill("solid", fgColor="FFF2CC")

        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(ncols)}{sheet.max_row}"
        side = Side(style="thin", color="999999")
        for col, heading in enumerate(layout["headings"], 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(55, max(14, len(heading) + 4))
        for row in sheet.iter_rows(min_row=3, max_row=sheet.max_row, max_col=ncols):
            for cell in row:
                cell.border = Border(left=side, right=side, top=side, bottom=side)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    if skipped:
        report = book.create_sheet("Skipped Pages")
        report.append(["PDF page", "Layout", "Reason"])
        for cell in report[1]:
            cell.font = Font(bold=True)
        for page, layout_name, reason in skipped:
            report.append([page, layout_name, reason])
        report.column_dimensions["A"].width = 12
        report.column_dimensions["B"].width = 28
        report.column_dimensions["C"].width = 85
        report.freeze_panes = "A2"
        report.auto_filter.ref = report.dimensions
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print("%s - %s" % (self.address_string(), fmt % args))

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/":
            body = HTML.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        if parsed.path != "/page":
            self.send_error(404)
            return
        q = parse_qs(parsed.query)
        try:
            session = SESSIONS[q.get("token", [""])[0]]
            page = int(q.get("n", ["1"])[0])
            samples, width, height = render(session, page)
            image = np.frombuffer(samples, dtype=np.uint8).reshape(height, width)
            ok, encoded = cv2.imencode(".png", image)
            if not ok:
                raise RuntimeError("Failed to encode rendered page.")
            body = encoded.tobytes()
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
        except (KeyError, ValueError, RuntimeError) as exc:
            json_reply(self, {"error": str(exc)}, 400)

    def do_POST(self):
        parsed = urlparse(self.path)
        token = parse_qs(parsed.query).get("token", [""])[0]
        if parsed.path == "/upload":
            self.upload()
        elif parsed.path == "/layout":
            self.save_layout(token)
        elif parsed.path == "/export":
            self.export(token)
        else:
            self.send_error(404)

    def body(self, limit=MAX_UPLOAD + 1024 * 1024):
        length = int(self.headers.get("Content-Length", "0"))
        if length < 1 or length > limit:
            raise ValueError("Request is empty or too large.")
        data = self.rfile.read(length)
        if len(data) != length:
            raise ValueError("Incomplete request.")
        return data

    def upload(self):
        path = None
        try:
            body = self.body()
            match = re.search(r'boundary=(?:"([^"]+)"|([^;]+))', self.headers.get("Content-Type", ""))
            if not match:
                raise ValueError("Expected a PDF upload.")
            boundary = (match.group(1) or match.group(2)).encode("ascii")
            payload, filename = None, "document.pdf"
            for part in body.split(b"--" + boundary):
                header, sep, contents = part.partition(b"\r\n\r\n")
                if sep and b'name="pdf"' in header:
                    fm = re.search(rb'filename="([^"]*)"', header)
                    if fm:
                        filename = Path(fm.group(1).decode("utf-8", "replace")).name
                    payload = contents.rsplit(b"\r\n", 1)[0]
                    break
            if not payload or not payload.startswith(b"%PDF-"):
                raise ValueError("Uploaded file is not a valid PDF.")
            if len(payload) > MAX_UPLOAD:
                raise ValueError("PDF exceeds the 150 MB limit.")
            token = uuid.uuid4().hex
            path = Path(tempfile.gettempdir()) / f"guided-pdf-{token}.pdf"
            path.write_bytes(payload)
            with fitz.open(path) as doc:
                pages = doc.page_count
            if pages < 1:
                raise ValueError("PDF has no pages.")
            SESSIONS[token] = {"path": path, "name": filename, "pages": pages, "layouts": []}
            json_reply(self, {"token": token, "name": filename, "pages": pages})
        except (ValueError, OSError, fitz.FileDataError) as exc:
            if path and path.exists():
                path.unlink(missing_ok=True)
            json_reply(self, {"error": str(exc)}, 400)

    def save_layout(self, token):
        try:
            session = SESSIONS[token]
            data = json.loads(self.body(1024 * 1024).decode("utf-8"))
            name = str(data.get("name", "")).strip()
            headings = [str(value).strip() for value in data.get("headings", [])]
            start, end = int(data["start"]), int(data["end"])
            rect = [float(value) for value in data["rect"]]
            dividers = [float(value) for value in data.get("dividers", [])]
            if not name or not headings or any(not h for h in headings):
                raise ValueError("Enter a layout name and at least one heading.")
            if len(dividers) > len(headings) - 1:
                raise ValueError("There are more dividers than column gaps.")
            if not 1 <= start <= end <= session["pages"]:
                raise ValueError("Page range is outside the PDF.")
            if (len(rect) != 4 or rect[0] < 0 or rect[1] < 0 or rect[2] <= 0 or
                    rect[3] <= 0 or rect[0] + rect[2] > 1.001 or rect[1] + rect[3] > 1.001):
                raise ValueError("Marked table region is invalid.")
            if any(not 0 < d < 1 for d in dividers) or any(
                b - a < 0.005 for a, b in zip([0.0] + dividers, dividers + [1.0])
            ):
                raise ValueError("Divider marks overlap or fall outside the table.")
            pages = set(range(start, end + 1))
            for old in session["layouts"]:
                if pages.intersection(range(old["start"], old["end"] + 1)):
                    raise ValueError(f"Page range overlaps layout '{old['name']}'.")
            session["layouts"].append({
                "name": name, "headings": headings, "start": start, "end": end,
                "rect": rect, "dividers": sorted(dividers)
            })
            json_reply(self, {"layouts": [
                {"name": item["name"], "start": item["start"], "end": item["end"]}
                for item in session["layouts"]
            ]})
        except KeyError:
            json_reply(self, {"error": "Upload the PDF before saving a layout."}, 400)
        except (ValueError, TypeError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            json_reply(self, {"error": str(exc)}, 400)

    def export(self, token):
        try:
            session = SESSIONS[token]
            content = create_workbook(session)
            filename = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(session["name"]).stem) + "_parts.xlsx"
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("X-Filename", filename)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        except KeyError:
            json_reply(self, {"error": "Upload the PDF before exporting."}, 400)
        except (ValueError, OSError, RuntimeError, pytesseract.TesseractNotFoundError) as exc:
            json_reply(self, {"error": str(exc)}, 400)
        except pytesseract.TesseractError as exc:
            json_reply(self, {"error": f"OCR failed: {exc}"}, 500)


def main():
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Guided PDF Table Extractor is running at http://{HOST}:{PORT}")
    print("Keep this terminal open. Press Ctrl+C to stop.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping.")
    finally:
        server.server_close()
        for session in SESSIONS.values():
            try:
                Path(session["path"]).unlink(missing_ok=True)
            except OSError as exc:
                print(f"Could not remove temporary upload: {exc}")


if __name__ == "__main__":
    main()
