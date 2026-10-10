import io
import json
import os
import re
import shutil
import tempfile
import threading
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

import cv2
import fitz
import numpy as np
import openpyxl
import pytesseract
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border


HOST = "127.0.0.1"
PORT = 8765
UPLOAD_LIMIT = 150 * 1024 * 1024
RENDER_DPI = 180
SESSIONS = {}
TESSERACT_CMD = os.environ.get("TESSERACT_CMD", "")
if not TESSERACT_CMD:
    default_tesseract = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
    if default_tesseract.exists():
        TESSERACT_CMD = str(default_tesseract)
    elif shutil.which("tesseract"):
        TESSERACT_CMD = shutil.which("tesseract")
if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


PAGE = r"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>PDF Table Extractor</title>
<style>
body{font:15px system-ui,Segoe UI,sans-serif;margin:0;background:#f3f5f8;color:#1c2530}
header{background:#18334f;color:white;padding:16px 24px;font-size:21px;font-weight:650}
main{max-width:1250px;margin:20px auto;padding:0 16px}
section{background:white;border:1px solid #d8dee7;border-radius:8px;padding:16px;margin:14px 0}
h2{font-size:18px;margin:0 0 12px}
.row{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin:8px 0}
button{background:#1769aa;color:white;border:0;border-radius:5px;padding:9px 13px;cursor:pointer}
button.secondary{background:#536579} button:disabled{opacity:.45;cursor:not-allowed}
input,select{padding:8px;border:1px solid #adb8c5;border-radius:4px}
input[type=text]{min-width:270px} input[type=number]{width:90px}
#preview{position:relative;display:inline-block;max-width:100%;line-height:0;cursor:crosshair}
#pageImage{display:block;max-width:100%;height:auto}
#overlay{position:absolute;inset:0;width:100%;height:100%;pointer-events:auto}
.muted{color:#647386;font-size:13px}
#status{white-space:pre-wrap;margin-top:10px}
#templates{padding-left:22px}.warn{color:#a33} code{background:#eef1f4;padding:2px 4px}
</style>
</head>
<body>
<header>PDF Table Extractor</header>
<main>
<section>
<h2>1. Upload a PDF</h2>
<div class="row"><input id="pdf" type="file" accept="application/pdf">
<button id="upload">Upload PDF</button></div><div id="status" class="muted">No PDF loaded.</div>
</section>
<section id="setup" hidden>
<h2>2. Mark a reference table</h2>
<p class="muted">Choose a representative parts-list page (not a drawing page). Drag around the parts rows, starting at the horizontal border just below the column-heading row; exclude the heading row. Then click internal column dividers from left to right. Enter one heading per column, separated by commas. Assign only pages that use this same layout.</p>
<div class="row"><button class="secondary" id="prev">Previous page</button>
<label>PDF page <input id="pageNo" type="number" min="1" value="1"></label>
<button class="secondary" id="next">Next page</button>
<span id="pageCount" class="muted"></span></div>
<div id="preview"><img id="pageImage" alt="PDF page preview"><canvas id="overlay"></canvas></div>
<div class="row"><button class="secondary" id="clearMarks">Clear marks</button>
<span id="markState" class="muted">Drag around the body rows, excluding the heading row.</span></div>
<div class="row"><label>Layout name <input id="layoutName" type="text" placeholder="e.g. Parts list A"></label>
<label>Column headings <input id="headings" type="text" placeholder="Item no, Qty, Designation"></label></div>
<div class="row"><label>Pages from <input id="pageStart" type="number" min="1"></label>
<label>through <input id="pageEnd" type="number" min="1"></label>
<button id="addTemplate">Save layout and page range</button></div>
<p class="muted">The marked geometry is applied only to pages in the range you assign. Other pages—including diagrams—are ignored. To handle a format change, save another layout and assign its page range.</p>
<ul id="templates"></ul>
<div class="row"><button id="export">Extract assigned pages and download Excel</button></div>
</section>
</main>
<script>
const $=id=>document.getElementById(id);
let token=null, pageCount=0, pageW=0,pageH=0, box=null, dividers=[], dragging=null, scaleX=1,scaleY=1;
function say(message,isError=false){$("status").textContent=message;$("status").className=isError?"warn":""}
function currentPage(){let n=Number($("pageNo").value);return Math.max(1,Math.min(pageCount,n||1))}
function drawMarks(){
 const img=$("pageImage"), c=$("overlay"); if(!img.naturalWidth)return;
 const r=img.getBoundingClientRect();
 const pixelWidth=Math.round(r.width*devicePixelRatio),pixelHeight=Math.round(r.height*devicePixelRatio);
 if(c.width!==pixelWidth)c.width=pixelWidth;if(c.height!==pixelHeight)c.height=pixelHeight;
 c.style.width=r.width+"px";c.style.height=r.height+"px";scaleX=c.width/pageW;scaleY=c.height/pageH;
 const ctx=c.getContext("2d");ctx.setTransform(scaleX,0,0,scaleY,0,0);ctx.clearRect(0,0,pageW,pageH);
 if(box){ctx.strokeStyle="#00a35a";ctx.lineWidth=Math.max(3,pageW/500);ctx.strokeRect(box.x,box.y,box.w,box.h);
  ctx.fillStyle="rgba(0,163,90,.12)";ctx.fillRect(box.x,box.y,box.w,box.h);
  ctx.strokeStyle="#e04725";ctx.fillStyle="#e04725";ctx.lineWidth=Math.max(2,pageW/700);
  dividers.forEach(x=>{ctx.beginPath();ctx.moveTo(x,box.y);ctx.lineTo(x,box.y+box.h);ctx.stroke()})}
 if(dragging){ctx.strokeStyle="#1769aa";ctx.lineWidth=Math.max(3,pageW/500);ctx.strokeRect(dragging.x,dragging.y,dragging.w,dragging.h)}
}
async function loadPage(n){
 if(!token)return; n=Math.max(1,Math.min(pageCount,n));$("pageNo").value=n;
 const im=$("pageImage");im.onload=()=>{pageW=im.naturalWidth;pageH=im.naturalHeight;drawMarks()};
 im.src=`/page?token=${encodeURIComponent(token)}&n=${n}&dpi=180&v=${Date.now()}`;
}
function point(e){const r=$("pageImage").getBoundingClientRect();return {
 x:Math.max(0,Math.min(pageW,(e.clientX-r.left)*pageW/r.width)),
 y:Math.max(0,Math.min(pageH,(e.clientY-r.top)*pageH/r.height))}}
const canvas=$("overlay");
canvas.addEventListener("pointerdown",e=>{if(!token)return;const p=point(e);
 if(box){const hit=dividers.some(x=>Math.abs(p.x-x)<pageW*.012&&p.y>=box.y&&p.y<=box.y+box.h);
  if(hit){const x=dividers.reduce((a,b)=>Math.abs(b-p.x)<Math.abs(a-p.x)?b:a);dividers=dividers.filter(d=>d!==x);drawMarks();$("markState").textContent="Divider removed. Click inside the table to add it back.";return}
  if(p.y>=box.y&&p.y<=box.y+box.h&&p.x>box.x&&p.x<box.x+box.w){dividers.push(p.x);dividers.sort((a,b)=>a-b);drawMarks();$("markState").textContent=`${dividers.length} divider(s) marked; click a red divider to remove it.`;return}}
 dragging={x:p.x,y:p.y,w:0,h:0};canvas.setPointerCapture(e.pointerId);drawMarks()});
canvas.addEventListener("pointermove",e=>{if(!dragging)return;const p=point(e);dragging.w=p.x-dragging.x;dragging.h=p.y-dragging.y;drawMarks()});
canvas.addEventListener("pointerup",()=>{if(!dragging)return;let b={...dragging};dragging=null;
 if(b.w<0){b.x+=b.w;b.w=-b.w}if(b.h<0){b.y+=b.h;b.h=-b.h}
 if(b.w>10&&b.h>10){box=b;dividers=[];$("markState").textContent="Body rows marked. Click internal column dividers from left to right."}drawMarks()});
$("upload").onclick=async()=>{const f=$("pdf").files[0];if(!f){say("Choose a PDF first.",true);return}
 if(f.size>150*1024*1024){say("PDF exceeds the 150 MB upload limit.",true);return}
 const fd=new FormData();fd.append("pdf",f);say("Uploading and reading PDF…");
 try{const r=await fetch("/upload",{method:"POST",body:fd});const d=await r.json();if(!r.ok)throw Error(d.error);
 token=d.token;pageCount=d.pages;$("pageCount").textContent=`of ${pageCount}`;
 $("pageStart").max=pageCount;$("pageEnd").max=pageCount;$("pageStart").value=1;$("pageEnd").value=pageCount;
 $("setup").hidden=false;box=null;dividers=[];say(`Loaded ${d.name} (${pageCount} pages). No page is treated as a table until you assign it to a marked layout.`);
 loadPage(1)}catch(e){say(e.message,true)}};
$("prev").onclick=()=>loadPage(currentPage()-1);$("next").onclick=()=>loadPage(currentPage()+1);
$("pageNo").onchange=()=>loadPage(currentPage());$("pageImage").onload=drawMarks;window.addEventListener("resize",drawMarks);
$("clearMarks").onclick=()=>{box=null;dividers=[];drawMarks();$("markState").textContent="Drag around the body rows, excluding the heading row."};
function showTemplates(items){$("templates").innerHTML=items.map(t=>`<li>${escapeHtml(t.name)} — pages ${t.start}–${t.end}; ${t.headings.length} columns (${escapeHtml(t.headings.join(" | "))})</li>`).join("")}
function escapeHtml(s){return s.replace(/[&<>"']/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;","'":"&#39;"}[c]))}
$("addTemplate").onclick=async()=>{if(!box){say("Mark the table region first.",true);return}
 const headings=$("headings").value.split(",").map(x=>x.trim()).filter(Boolean);
 const start=Number($("pageStart").value),end=Number($("pageEnd").value),name=$("layoutName").value.trim();
 if(!name||headings.length<2||dividers.length!==headings.length-1){say(`Enter a layout name, at least two headings, and exactly ${Math.max(0,headings.length-1)} internal divider(s).`,true);return}
 if(!start||!end||start>end||start<1||end>pageCount){say("Enter a valid inclusive page range.",true);return}
 const list=[...dividers].sort((a,b)=>a-b),all=[box.x,...list,box.x+box.w];
 if(all.some((x,i)=>i>0&&x-all[i-1]<pageW*.005)){say("Column boundaries overlap or are too close; remark the dividers.",true);return}
 const item={name,headings,start,end,rect:[box.x/pageW,box.y/pageH,box.w/pageW,box.h/pageH],edges:all.map(x=>(x-box.x)/box.w)};
 try{const r=await fetch("/template?token="+encodeURIComponent(token),{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(item)});
 const d=await r.json();if(!r.ok)throw Error(d.error);showTemplates(d.templates);say("Layout saved. Add another layout if the table format changes.")}catch(e){say(e.message,true)}};
$("export").onclick=async()=>{if(!token){say("Upload a PDF first.",true);return}
 $("export").disabled=true;say("Extracting assigned pages with OCR…");
 try{const r=await fetch("/export?token="+encodeURIComponent(token),{method:"POST"});if(!r.ok){const d=await r.json();throw Error(d.error)}
 const blob=await r.blob(),url=URL.createObjectURL(blob),a=document.createElement("a");a.href=url;a.download=r.headers.get("X-Filename")||"extracted_parts.xlsx";a.click();URL.revokeObjectURL(url);
 say("Excel workbook downloaded. Pages outside your assigned ranges were not processed.")}catch(e){say(e.message,true)}finally{$("export").disabled=false}};
</script>
</body></html>"""


def json_response(handler, payload, status=200):
    body = json.dumps(payload).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def render_page(session, page_number, dpi=RENDER_DPI):
    with fitz.open(session["pdf"]) as doc:
        if page_number < 1 or page_number > doc.page_count:
            raise ValueError("Page number is outside this PDF.")
        pix = doc[page_number - 1].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY, alpha=False)
        return pix.samples, pix.width, pix.height


def horizontal_rules(gray):
    height, width = gray.shape
    ink = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    kernel_width = max(20, int(width * 0.45))
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (kernel_width, 1))
    lines = cv2.morphologyEx(ink, cv2.MORPH_OPEN, kernel)
    strength = np.count_nonzero(lines, axis=1)
    ys = np.flatnonzero(strength >= width * 0.45)
    if not len(ys):
        return []
    gap = max(3, int(height * 0.004))
    groups = np.split(ys, np.where(np.diff(ys) > gap)[0] + 1)
    return [int(round(float(np.mean(g)))) for g in groups if len(g)]


def extract_page(session, page_number, template):
    samples, width, height = render_page(session, page_number)
    gray = np.frombuffer(samples, dtype=np.uint8).reshape(height, width).copy()
    rx, ry, rw, rh = template["rect"]
    x0, y0 = int(round(rx * width)), int(round(ry * height))
    x1, y1 = int(round((rx + rw) * width)), int(round((ry + rh) * height))
    x0, x1 = max(0, x0), min(width, x1)
    y0, y1 = max(0, y0), min(height, y1)
    if x1 - x0 < 20 or y1 - y0 < 20:
        raise ValueError(f"Page {page_number}: marked table area is too small.")
    local = gray[y0:y1, x0:x1]
    rules = horizontal_rules(local)
    min_rules = 2
    if len(rules) < min_rules:
        raise ValueError(
            f"Page {page_number}: found only {len(rules)} horizontal row boundaries "
            "inside the assigned table area. Check its layout or that page's assignment."
        )
    edges = [int(round(x0 + p * (x1 - x0))) for p in template["edges"]]
    rows = []
    bands = list(zip(rules, rules[1:]))
    for top, bottom in bands:
        if bottom - top < max(5, int(height * 0.003)):
            continue
        row = []
        for left, right in zip(edges, edges[1:]):
            pad_x = max(2, int(width * 0.002))
            pad_y = max(2, int(height * 0.0015))
            cell = gray[min(y1, y0 + top + pad_y):max(y0, y0 + bottom - pad_y),
                        min(x1, left + pad_x):max(x0, right - pad_x)]
            if cell.size == 0 or np.count_nonzero(cell < 170) < 5:
                row.append(("", 100))
                continue
            enlarged = cv2.resize(cell, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC)
            data = pytesseract.image_to_data(
                enlarged, config="--psm 6", lang="eng",
                output_type=pytesseract.Output.DICT
            )
            words = [(text.strip(), float(conf)) for text, conf in zip(data["text"], data["conf"])
                     if text.strip() and float(conf) >= 0]
            if words:
                text = re.sub(r"\s+", " ", " ".join(word for word, _ in words))
                confidence = sum(conf for _, conf in words) / len(words)
                row.append((text, confidence))
            else:
                row.append(("?", 0))
        if any(text for text, _ in row):
            rows.append(row)
    return rows


def make_workbook(session):
    templates = session["templates"]
    if not templates:
        raise ValueError("Save at least one reference layout and page range before exporting.")
    ranges = []
    for template in templates:
        for page in range(template["start"], template["end"] + 1):
            ranges.append((page, template))
    ranges.sort(key=lambda pair: pair[0])
    seen = set()
    for page, _ in ranges:
        if page in seen:
            raise ValueError(f"Page {page} is assigned to more than one layout.")
        seen.add(page)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names = set()
    for template in templates:
        raw_name = re.sub(r"[\[\]:*?/\\]", "-", template["name"]).strip() or "Parts"
        base = raw_name[:31]
        name, suffix = base, 2
        while name.casefold() in used_names:
            tail = f" ({suffix})"
            name, suffix = base[:31 - len(tail)] + tail, suffix + 1
        used_names.add(name.casefold())
        ws = wb.create_sheet(name)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(template["headings"]))
        ws.cell(row=1, column=1, value=template["name"]).font = Font(bold=True, size=14)
        ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
        for column, heading in enumerate(template["headings"], start=1):
            ws.cell(row=3, column=column, value=heading)
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        ws.row_dimensions[1].height = 24
        pages = range(template["start"], template["end"] + 1)
        count = 0
        for page in pages:
            for row in extract_page(session, page, template):
                values = (row + [("", 100)] * len(template["headings"]))[:len(template["headings"])]
                ws.append([text for text, _ in values])
                for column, (text, confidence) in enumerate(values, start=1):
                    if text and confidence < 70:
                        ws.cell(row=ws.max_row, column=column).fill = PatternFill(
                            "solid", fgColor="FFF2CC"
                        )
                count += 1
        if not count:
            raise ValueError(f"Layout '{template['name']}' produced no non-empty rows.")
        ws.freeze_panes = "A4"
        ws.auto_filter.ref = f"A3:{openpyxl.utils.get_column_letter(len(template['headings']))}{ws.max_row}"
        for col, heading in enumerate(template["headings"], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(
                55, max(14, len(heading) + 4)
            )
        side = Side(style="thin", color="999999")
        for row in ws.iter_rows(min_row=3, max_row=ws.max_row,
                                min_col=1, max_col=len(template["headings"])):
            for cell in row:
                cell.border = Border(left=side, right=side, top=side, bottom=side)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print("%s - %s" % (self.address_string(), fmt % args))

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/":
            body = PAGE.encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "text/html; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)
            return
        from urllib.parse import parse_qs
        params = parse_qs(parsed.query)
        token = params.get("token", [""])[0]
        if parsed.path == "/page":
            try:
                session = SESSIONS[token]
                page_number = int(params.get("n", ["1"])[0])
                dpi = int(params.get("dpi", [str(RENDER_DPI)])[0])
                if dpi < 72 or dpi > 300:
                    raise ValueError("Preview DPI must be from 72 to 300.")
                samples, width, height = render_page(session, page_number, dpi)
                image = np.frombuffer(samples, dtype=np.uint8).reshape(height, width)
                ok, encoded = cv2.imencode(".png", image)
                if not ok:
                    raise RuntimeError("Could not encode the rendered page preview.")
                body = encoded.tobytes()
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            except (KeyError, ValueError, RuntimeError) as exc:
                json_response(self, {"error": str(exc)}, 400)
            return
        self.send_error(404)

    def do_POST(self):
        parsed = urlparse(self.path)
        from urllib.parse import parse_qs
        token = parse_qs(parsed.query).get("token", [""])[0]
        if parsed.path == "/upload":
            self.upload()
        elif parsed.path == "/template":
            self.save_template(token)
        elif parsed.path == "/export":
            self.export(token)
        else:
            self.send_error(404)

    def read_body(self):
        length = int(self.headers.get("Content-Length", "0"))
        if length <= 0 or length > UPLOAD_LIMIT + 1024 * 1024:
            raise ValueError("Request is empty or exceeds the upload limit.")
        return self.rfile.read(length)

    def upload(self):
        try:
            body = self.read_body()
            content_type = self.headers.get("Content-Type", "")
            match = re.search(r"boundary=(?:\"([^\"]+)\"|([^;]+))", content_type)
            if not match:
                raise ValueError("Expected a PDF file upload.")
            boundary = (match.group(1) or match.group(2)).encode("ascii")
            parts = body.split(b"--" + boundary)
            payload, filename = None, "document.pdf"
            for part in parts:
                head, sep, value = part.partition(b"\r\n\r\n")
                if not sep or b'name="pdf"' not in head:
                    continue
                name_match = re.search(rb'filename="([^"]*)"', head)
                if name_match:
                    filename = Path(name_match.group(1).decode("utf-8", "replace")).name
                payload = value.rsplit(b"\r\n", 1)[0]
                break
            if not payload:
                raise ValueError("No PDF was found in the upload.")
            if not payload.startswith(b"%PDF-"):
                raise ValueError("The uploaded file is not a valid PDF.")
            if len(payload) > UPLOAD_LIMIT:
                raise ValueError("PDF exceeds the 150 MB upload limit.")
            token = uuid.uuid4().hex
            path = Path(tempfile.gettempdir()) / f"pdf-table-{token}.pdf"
            path.write_bytes(payload)
            try:
                with fitz.open(path) as doc:
                    if not doc.page_count:
                        raise ValueError("PDF has no pages.")
                    pages = doc.page_count
            except Exception:
                path.unlink(missing_ok=True)
                raise
            SESSIONS[token] = {"pdf": path, "pages": pages, "name": filename, "templates": []}
            json_response(self, {"token": token, "pages": pages, "name": filename})
        except (ValueError, OSError, fitz.FileDataError) as exc:
            json_response(self, {"error": str(exc)}, 400)

    def save_template(self, token):
        try:
            session = SESSIONS[token]
            data = json.loads(self.read_body().decode("utf-8"))
            name = str(data.get("name", "")).strip()
            headings = [str(h).strip() for h in data.get("headings", [])]
            start, end = int(data["start"]), int(data["end"])
            rect = [float(v) for v in data["rect"]]
            edges = [float(v) for v in data["edges"]]
            if not name or not headings or any(not h for h in headings):
                raise ValueError("Layout name and non-empty headings are required.")
            if len(edges) != len(headings) + 1 or edges[0] < -0.0001 or edges[-1] > 1.0001:
                raise ValueError("Each heading needs a corresponding column boundary.")
            if any(b <= a for a, b in zip(edges, edges[1:])):
                raise ValueError("Column boundaries must be strictly left-to-right.")
            if not (1 <= start <= end <= session["pages"]):
                raise ValueError("Page range is outside the uploaded PDF.")
            if len(rect) != 4 or rect[2] <= 0 or rect[3] <= 0 or rect[0] < 0 or rect[1] < 0 or rect[0] + rect[2] > 1.001 or rect[1] + rect[3] > 1.001:
                raise ValueError("Marked table rectangle is invalid.")
            overlap = set(range(start, end + 1))
            for old in session["templates"]:
                if overlap.intersection(range(old["start"], old["end"] + 1)):
                    raise ValueError(
                        f"Page ranges overlap with layout '{old['name']}'. "
                        "Each PDF page can use only one layout."
                    )
            template = {"name": name, "headings": headings, "start": start, "end": end,
                        "rect": rect, "edges": edges}
            session["templates"].append(template)
            json_response(self, {"templates": [
                {"name": t["name"], "headings": t["headings"], "start": t["start"], "end": t["end"]}
                for t in session["templates"]
            ]})
        except KeyError:
            json_response(self, {"error": "Upload a PDF before saving a layout."}, 400)
        except (ValueError, TypeError, json.JSONDecodeError, UnicodeDecodeError) as exc:
            json_response(self, {"error": str(exc)}, 400)

    def export(self, token):
        try:
            session = SESSIONS[token]
            content = make_workbook(session).getvalue()
            filename = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(session["name"]).stem) + "_parts.xlsx"
            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
            self.send_header("X-Filename", filename)
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        except KeyError:
            json_response(self, {"error": "Upload a PDF before exporting."}, 400)
        except (ValueError, OSError, RuntimeError, pytesseract.TesseractNotFoundError) as exc:
            json_response(self, {"error": str(exc)}, 400)
        except pytesseract.TesseractError as exc:
            json_response(self, {"error": f"OCR failed: {exc}"}, 500)


def main():
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"PDF Table Extractor is ready at http://{HOST}:{PORT}")
    print("This app listens only on this computer. Press Ctrl+C to stop.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping server…")
    finally:
        server.server_close()
        for session in SESSIONS.values():
            try:
                Path(session["pdf"]).unlink(missing_ok=True)
            except OSError as exc:
                print(f"Could not remove temporary upload {session['pdf']}: {exc}")


if __name__ == "__main__":
    main()
